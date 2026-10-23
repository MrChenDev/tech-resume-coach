#!/usr/bin/env python3
"""Create an editable, single-column DOCX and optionally a PDF from resume JSON."""

from __future__ import annotations

import argparse
import html
import json
import os
import sys
from pathlib import Path
from typing import Any


def add_heading(document: Any, text: str, level: int = 1) -> None:
    if text:
        document.add_heading(str(text), level=level)


def add_bullets(document: Any, items: Any) -> None:
    for item in items or []:
        if isinstance(item, dict):
            text = item.get("text", "")
        else:
            text = str(item)
        if text.strip():
            document.add_paragraph(text.strip(), style="List Bullet")


def add_records(document: Any, title: str, records: Any) -> None:
    if not records:
        return
    add_heading(document, title, 1)
    for record in records:
        if isinstance(record, str):
            document.add_paragraph(record)
            continue
        lead_parts = [record.get(k, "") for k in ("title", "organization", "role")]
        lead = " | ".join(str(value).strip() for value in lead_parts if str(value).strip())
        dates = str(record.get("dates", "")).strip()
        if lead or dates:
            paragraph = document.add_paragraph()
            if lead:
                paragraph.add_run(lead).bold = True
            if dates:
                paragraph.add_run(f"  {dates}")
        if record.get("description"):
            document.add_paragraph(str(record["description"]))
        add_bullets(document, record.get("bullets"))


def build_docx(data: dict[str, Any], output: Path) -> None:
    from docx import Document
    from docx.shared import Inches, Pt

    document = Document()
    section = document.sections[0]
    section.top_margin = Inches(0.55)
    section.bottom_margin = Inches(0.55)
    section.left_margin = Inches(0.65)
    section.right_margin = Inches(0.65)
    normal = document.styles["Normal"]
    normal.font.name = "Aptos"
    normal.font.size = Pt(10)

    name = str(data.get("name", "")).strip()
    if name:
        paragraph = document.add_paragraph()
        paragraph.alignment = 1
        run = paragraph.add_run(name)
        run.bold = True
        run.font.size = Pt(18)
    contact = " | ".join(str(value).strip() for value in data.get("contact", []) if str(value).strip())
    if contact:
        paragraph = document.add_paragraph(contact)
        paragraph.alignment = 1
    if data.get("objective"):
        document.add_paragraph(str(data["objective"]))
    if data.get("summary"):
        add_heading(document, "个人简介", 1)
        document.add_paragraph(str(data["summary"]))

    add_records(document, "教育经历", data.get("education"))
    add_bullets_section = data.get("skills")
    if add_bullets_section:
        add_heading(document, "专业技能", 1)
        add_bullets(document, add_bullets_section)
    add_records(document, "实习经历", data.get("internships"))
    add_records(document, "项目经历", data.get("projects"))
    add_records(document, "科研 / 竞赛 / 开源", data.get("activities"))
    add_records(document, "荣誉与证书", data.get("awards"))

    output.parent.mkdir(parents=True, exist_ok=True)
    document.save(str(output))


def find_cjk_font() -> Path:
    configured = os.environ.get("RESUME_CJK_FONT")
    candidates = [Path(configured)] if configured else []
    candidates.extend(
        Path(value)
        for value in (
            r"C:\Windows\Fonts\simhei.ttf",
            r"C:\Windows\Fonts\msyh.ttc",
            r"C:\Windows\Fonts\simsun.ttc",
            "/System/Library/Fonts/PingFang.ttc",
            "/System/Library/Fonts/STHeiti Light.ttc",
            "/usr/share/fonts/opentype/noto/NotoSansCJK-Regular.ttc",
            "/usr/share/fonts/truetype/noto/NotoSansCJK-Regular.ttc",
            "/usr/share/fonts/truetype/wqy/wqy-zenhei.ttc",
        )
    )
    for candidate in candidates:
        if candidate.is_file():
            return candidate
    raise FileNotFoundError(
        "No Chinese font found. Install a CJK font or set RESUME_CJK_FONT to a .ttf/.ttc font file."
    )


def build_pdf(data: dict[str, Any], output: Path) -> None:
    from reportlab.lib import colors
    from reportlab.lib.enums import TA_CENTER
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
    from reportlab.pdfbase import pdfmetrics
    from reportlab.pdfbase.ttfonts import TTFont
    from reportlab.platypus import KeepTogether, Paragraph, SimpleDocTemplate, Spacer

    font_path = find_cjk_font()
    pdfmetrics.registerFont(TTFont("ResumeCJK", str(font_path), subfontIndex=0))
    styles = getSampleStyleSheet()
    title_style = ParagraphStyle(
        "ResumeTitle", parent=styles["Title"], fontName="ResumeCJK", fontSize=18,
        leading=23, alignment=TA_CENTER, textColor=colors.HexColor("#202124"), spaceAfter=5,
    )
    contact_style = ParagraphStyle(
        "ResumeContact", parent=styles["Normal"], fontName="ResumeCJK", fontSize=9,
        leading=13, alignment=TA_CENTER, textColor=colors.HexColor("#444444"), spaceAfter=7,
    )
    body_style = ParagraphStyle(
        "ResumeBody", parent=styles["BodyText"], fontName="ResumeCJK", fontSize=9.5,
        leading=14, textColor=colors.HexColor("#222222"), spaceAfter=3,
    )
    heading_style = ParagraphStyle(
        "ResumeSection", parent=styles["Heading2"], fontName="ResumeCJK", fontSize=11.5,
        leading=15, textColor=colors.HexColor("#15385b"), spaceBefore=7, spaceAfter=3,
        keepWithNext=True,
    )

    def paragraph_text(value: Any) -> str:
        return html.escape(str(value)).replace("\n", "<br/>")

    story: list[Any] = []
    name = str(data.get("name", "")).strip()
    if name:
        story.append(Paragraph(paragraph_text(name), title_style))
    contacts = data.get("contact", [])
    if isinstance(contacts, str):
        contacts = [contacts]
    contact = " | ".join(str(value).strip() for value in contacts if str(value).strip())
    if contact:
        story.append(Paragraph(paragraph_text(contact), contact_style))
    if data.get("objective"):
        story.append(Paragraph(paragraph_text(data["objective"]), body_style))
    if data.get("summary"):
        story.extend([Paragraph("个人简介", heading_style), Paragraph(paragraph_text(data["summary"]), body_style)])

    def add_records(title: str, records: Any) -> None:
        if not records:
            return
        story.append(Paragraph(paragraph_text(title), heading_style))
        for record in records:
            if isinstance(record, str):
                story.append(Paragraph(paragraph_text(record), body_style))
                continue
            lead_parts = [record.get(k, "") for k in ("title", "organization", "role")]
            lead = " | ".join(str(value).strip() for value in lead_parts if str(value).strip())
            dates = str(record.get("dates", "")).strip()
            group: list[Any] = []
            if lead or dates:
                heading = "<b>" + paragraph_text(lead) + "</b>" if lead else ""
                if dates:
                    heading += ("  " if heading else "") + paragraph_text(dates)
                group.append(Paragraph(heading, body_style))
            if record.get("description"):
                group.append(Paragraph(paragraph_text(record["description"]), body_style))
            for bullet in record.get("bullets", []) or []:
                value = bullet.get("text", "") if isinstance(bullet, dict) else bullet
                if str(value).strip():
                    group.append(Paragraph("• " + paragraph_text(value), body_style))
            if group:
                story.append(KeepTogether(group))

    add_records("教育经历", data.get("education"))
    if data.get("skills"):
        story.append(Paragraph("专业技能", heading_style))
        for item in data["skills"]:
            value = item.get("text", "") if isinstance(item, dict) else item
            if str(value).strip():
                story.append(Paragraph("• " + paragraph_text(value), body_style))
    add_records("实习经历", data.get("internships"))
    add_records("项目经历", data.get("projects"))
    add_records("科研 / 竞赛 / 开源", data.get("activities"))
    add_records("荣誉与证书", data.get("awards"))

    output.parent.mkdir(parents=True, exist_ok=True)
    document = SimpleDocTemplate(
        str(output), pagesize=A4, rightMargin=42, leftMargin=42, topMargin=38, bottomMargin=38,
        title=name or "Resume", author="Tech Resume Coach",
    )
    document.build(story or [Spacer(1, 1)])


def main() -> int:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--input", required=True, type=Path, help="UTF-8 JSON following examples/resume.example.json")
    parser.add_argument("--output-dir", required=True, type=Path)
    parser.add_argument("--formats", default="docx,pdf", help="Comma-separated: docx,pdf (default: both)")
    args = parser.parse_args()

    formats = {part.strip().lower() for part in args.formats.split(",") if part.strip()}
    if not formats or not formats <= {"docx", "pdf"}:
        parser.error("--formats accepts docx and/or pdf")
    if not args.input.is_file():
        parser.error(f"Input JSON does not exist: {args.input}")

    try:
        data = json.loads(args.input.read_text(encoding="utf-8"))
        if not isinstance(data, dict):
            parser.error("The JSON root must be an object")
        args.output_dir.mkdir(parents=True, exist_ok=True)
        docx_path = args.output_dir / f"{args.input.stem}.docx"
        if "docx" in formats:
            build_docx(data, docx_path)
            print(f"Created: {docx_path}")
        if "pdf" in formats:
            pdf_path = args.output_dir / f"{args.input.stem}.pdf"
            build_pdf(data, pdf_path)
            print(f"Created: {pdf_path}")
    except ImportError as exc:
        print(f"Missing optional dependency: {exc}. Install packages listed in requirements.txt.", file=sys.stderr)
        return 2
    except (OSError, ValueError, RuntimeError) as exc:
        print(f"Export incomplete: {exc}", file=sys.stderr)
        return 1
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
